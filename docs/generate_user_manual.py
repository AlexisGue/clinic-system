#!/usr/bin/env python3
"""Generate Inventory System user manual (.docx)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\xampp\htdocs\inventory-system\docs\Manual_Usuario_Inventory_System.docx")


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, color=RGBColor(0x0F, 0x76, 0x6E))
    return p


def add_para(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)
        p.paragraph_format.space_after = Pt(2)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            set_run_font(run)
        p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10, color=RGBColor(0x0F, 0x76, 0x6E))
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("INVENTORY SYSTEM")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x0F, 0x76, 0x6E))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Manual de Usuario")
    set_run_font(r, size=20, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Guía práctica para operar el sistema de inventario, compras y ventas (POS)")
    set_run_font(r, size=12)
    r.italic = True

    ver = doc.add_paragraph()
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ver.add_run("Versión 1.0  ·  Julio 2026  ·  Moneda: USD")
    set_run_font(r, size=11, color=RGBColor(0x64, 0x74, 0x8B))

    doc.add_page_break()

    # TOC-like overview
    add_heading_styled(doc, "Índice", 1)
    add_bullets(doc, [
        "1. Introducción y acceso",
        "2. Roles y permisos",
        "3. Dashboard",
        "4. Catálogos (categorías, marcas, unidades, impuestos)",
        "5. Productos e inventario (kardex / ajustes)",
        "6. Proveedores y compras",
        "7. Clientes",
        "8. Punto de venta (POS) y caja",
        "9. Ventas",
        "10. Reportes",
        "11. Configuración de empresa",
        "12. Usuarios, roles y auditoría",
        "13. Flujos recomendados del día a día",
        "14. Preguntas frecuentes y problemas comunes",
    ])

    # 1
    add_heading_styled(doc, "1. Introducción y acceso", 1)
    add_para(doc, "Inventory System es una aplicación web para controlar inventario, compras, clientes y ventas en punto de caja (POS). La interfaz está en español; los montos se muestran en dólares (USD).")

    add_heading_styled(doc, "1.1 Cómo entrar", 2)
    add_steps(doc, [
        "Abre el navegador e ingresa a: http://localhost:5173",
        "Escribe tu correo y contraseña.",
        "Haz clic en «Entrar al sistema».",
    ])
    add_para(doc, "Credenciales de demostración (usuario administrador):", bold=True)
    add_table(doc, ["Campo", "Valor"], [
        ["Correo", "admin@demo.test"],
        ["Contraseña", "Admin1234"],
    ])
    add_para(doc, "Si olvidaste la contraseña o el usuario está inactivo, contacta al administrador del sistema. Tras varios intentos fallidos, el acceso se bloquea temporalmente por seguridad.")

    add_heading_styled(doc, "1.2 Pantalla principal", 2)
    add_bullets(doc, [
        "Menú izquierdo: módulos disponibles según tu rol.",
        "Encabezado superior: nombre del usuario y botón «Salir».",
        "Área central: contenido del módulo seleccionado.",
    ])

    # 2
    add_heading_styled(doc, "2. Roles y permisos", 1)
    add_para(doc, "Lo que ves y puedes hacer depende del rol asignado. Roles incluidos en la demo:")
    add_table(doc, ["Rol", "Uso típico", "Puede (resumen)"], [
        ["admin", "Administrador", "Todo: usuarios, roles, configuración, cancelar ventas/compras, reportes"],
        ["manager", "Gerente / encargado", "Operación completa de catálogos, productos, compras, ventas y reportes; ve settings pero no siempre edita"],
        ["cashier", "Cajero", "POS, ventas, clientes, abrir/cerrar caja; no cancela ventas ni administra usuarios"],
        ["warehouse", "Almacén", "Productos, inventario, compras y recepción"],
    ])
    add_para(doc, "Si un botón no aparece (por ejemplo «Cancelar venta»), suele ser por falta de permiso o porque la regla de negocio no lo permite (venta ya cancelada, compra ya recibida, etc.).")

    # 3
    add_heading_styled(doc, "3. Dashboard", 1)
    add_para(doc, "Menú: Dashboard. Muestra un resumen operativo del negocio.")
    add_bullets(doc, [
        "KPIs: ventas de hoy, ventas del periodo, compras del periodo, productos bajo stock.",
        "Gráfico de ventas de los últimos 7 días.",
        "Top productos por ingreso.",
        "Lista de alertas de stock bajo (atajo a Productos filtrados).",
    ])
    add_steps(doc, [
        "Elige rango de fechas (desde / hasta).",
        "Haz clic en «Actualizar».",
        "Opcional: entra a «Reportes» para análisis más detallados.",
    ])

    # 4
    add_heading_styled(doc, "4. Catálogos", 1)
    add_para(doc, "Antes de crear muchos productos, conviene tener listos los catálogos base.")

    add_heading_styled(doc, "4.1 Categorías, Marcas y Unidades", 2)
    add_steps(doc, [
        "Entra al menú correspondiente (Categorías / Marcas / Unidades).",
        "Usa «Nuevo…» para abrir el formulario.",
        "Completa nombre (y abreviación en Unidades, ej. pza, kg).",
        "Guarda. Para editar o eliminar usa las acciones de cada fila.",
        "«Cancelar» cierra el formulario sin guardar.",
    ])

    add_heading_styled(doc, "4.2 Impuestos", 2)
    add_para(doc, "Menú: Impuestos. Define tasas (ej. Sales Tax 7% o Exento 0%). Luego se asocian a cada producto y se calculan en la venta.")
    add_bullets(doc, [
        "Nombre: cómo se verá en listas.",
        "Tasa (%): porcentaje aplicado sobre el precio sin impuesto.",
        "Activo: si está desactivado no debería usarse en productos nuevos.",
    ])

    # 5
    add_heading_styled(doc, "5. Productos e inventario", 1)

    add_heading_styled(doc, "5.1 Crear o editar un producto", 2)
    add_steps(doc, [
        "Ve a Productos → «Nuevo producto».",
        "Completa SKU (único), nombre, unidad, precios de costo y venta.",
        "Asigna categoría, marca, impuesto (opcional pero recomendado).",
        "Define stock inicial (solo al crear) y stock mínimo para alertas.",
        "Puedes subir una imagen; el sistema la procesa de forma segura.",
        "Guarda. Para editar: acción «Editar» en la lista (el stock no se cambia desde el formulario de edición).",
    ])
    add_para(doc, "Importante: el stock no se edita a mano en el producto. Se mueve con compras, ventas o ajustes de inventario.", italic=True)

    add_heading_styled(doc, "5.2 Buscar y filtrar", 2)
    add_bullets(doc, [
        "Busca por nombre, SKU o código de barras.",
        "Filtra por categoría, marca o «bajo stock».",
        "Desde el Dashboard, «Ver todos» en alertas abre productos con stock bajo.",
    ])

    add_heading_styled(doc, "5.3 Kardex (historial de movimientos)", 2)
    add_steps(doc, [
        "En Productos, haz clic en «Kardex» del producto.",
        "Verás entradas/salidas: compra, venta, ajuste, devolución, inicial, etc.",
        "Cada movimiento muestra cantidad, stock antes/después, usuario y fecha.",
    ])
    add_para(doc, "El kardex es la fuente de verdad del inventario; el campo «stock» del producto es el saldo actual.", italic=True)

    add_heading_styled(doc, "5.4 Ajustar stock", 2)
    add_steps(doc, [
        "En la lista de productos, usa «Ajustar» (requiere permiso inventory.adjust).",
        "Elige entrada o salida, cantidad y motivo/nota.",
        "Confirma. El sistema valida que no quede stock negativo en salidas.",
    ])

    # 6
    add_heading_styled(doc, "6. Proveedores y compras", 1)

    add_heading_styled(doc, "6.1 Proveedores", 2)
    add_para(doc, "Menú: Proveedores. Alta de datos de contacto y Tax ID. Se usan al crear órdenes de compra.")

    add_heading_styled(doc, "6.2 Crear una compra", 2)
    add_steps(doc, [
        "Ve a Compras → «Nueva compra».",
        "Selecciona proveedor y fecha.",
        "Busca productos y agrégalos con cantidad y costo unitario.",
        "Ajusta impuestos de la compra si aplica.",
        "Guarda. Se genera un folio automático (ej. C-000001) con estado «Ordenada».",
    ])

    add_heading_styled(doc, "6.3 Recibir mercancía", 2)
    add_para(doc, "El stock NO aumenta al crear la compra; aumenta al recibir.")
    add_steps(doc, [
        "Abre la compra (Ver / Recibir).",
        "Indica cantidades recibidas por línea (puede ser recepción parcial).",
        "Confirma. El estado pasa a «Parcial» o «Recibida».",
        "Cada recepción genera movimientos de inventario tipo compra.",
    ])

    add_heading_styled(doc, "6.4 Cancelar compra", 2)
    add_bullets(doc, [
        "Solo si está «Ordenada» y aún no hay recepciones.",
        "Si ya recibiste mercancía, no se puede cancelar (regla de negocio).",
        "Requiere permiso purchases.cancel.",
    ])

    # 7
    add_heading_styled(doc, "7. Clientes", 1)
    add_para(doc, "Menú: Clientes. Úsalos en el POS para asociar la venta (opcional). Incluye Tax ID, teléfono, ciudad, etc.")
    add_steps(doc, [
        "«Nuevo cliente» → completa datos → Guardar.",
        "Edita o elimina desde la tabla según permisos.",
    ])
    add_para(doc, "Existe el cliente demo «Público en general» para ventas de mostrador sin datos fiscales.")

    # 8
    add_heading_styled(doc, "8. Punto de venta (POS) y caja", 1)
    add_para(doc, "Menú: Punto de venta. Es el módulo diario del cajero.")

    add_heading_styled(doc, "8.1 Abrir caja (obligatorio para vender)", 2)
    add_steps(doc, [
        "Entra a Punto de venta.",
        "Si no hay sesión abierta, indica el monto de apertura (efectivo inicial).",
        "Haz clic en «Abrir caja».",
        "Solo puede haber una caja abierta a la vez en el sistema.",
    ])
    add_para(doc, "Sin caja abierta, el sistema bloquea las ventas con el mensaje de sesión de caja requerida.", italic=True)

    add_heading_styled(doc, "8.2 Cobrar una venta", 2)
    add_steps(doc, [
        "Busca el producto (nombre / SKU) y agrégalo al carrito.",
        "Ajusta cantidades. Revisa subtotal, impuesto y total (USD).",
        "Opcional: selecciona cliente.",
        "Elige método de pago (efectivo, tarjeta, etc.) y confirma el monto (debe coincidir con el total).",
        "Haz clic en cobrar / registrar venta.",
        "Se genera folio V-000001, se descuenta stock y queda registrada la venta.",
    ])
    add_bullets(doc, [
        "Si no hay stock suficiente, la venta se rechaza.",
        "Usa «Vaciar carrito» para limpiar la venta en curso.",
    ])

    add_heading_styled(doc, "8.3 Cerrar caja", 2)
    add_steps(doc, [
        "Al final del turno, captura el monto de cierre contado.",
        "Haz clic en «Cerrar caja».",
        "Ya no podrás vender hasta abrir una nueva sesión.",
    ])

    # 9
    add_heading_styled(doc, "9. Ventas", 1)
    add_para(doc, "Menú: Ventas. Consulta el historial de tickets.")
    add_bullets(doc, [
        "Filtra por fechas o busca por folio.",
        "«Ver» abre el detalle: líneas, impuestos, pagos, usuario.",
        "«Cancelar» (si tienes permiso y la venta está completada) revierte stock y marca la venta como cancelada. No se borran registros.",
    ])
    add_para(doc, "El rol cajero normalmente no puede cancelar ventas; debe hacerlo un admin/manager.", italic=True)

    # 10
    add_heading_styled(doc, "10. Reportes", 1)
    add_para(doc, "Menú: Reportes (permiso reports.view / reports.export).")
    add_bullets(doc, [
        "Ventas: totales, impuestos y detalle por periodo.",
        "Compras: totales del periodo.",
        "Inventario: existencias, costos y valor de inventario.",
    ])
    add_steps(doc, [
        "Elige el tipo de reporte y el rango de fechas (cuando aplique).",
        "Consulta en pantalla o exporta a CSV / PDF.",
        "Los PDF muestran datos de la empresa (razón social, Tax ID, moneda USD).",
    ])

    # 11
    add_heading_styled(doc, "11. Configuración de empresa", 1)
    add_para(doc, "Menú: Configuración. Solo quien tenga settings.update (admin) puede guardar cambios.")
    add_bullets(doc, [
        "Razón social / nombre comercial",
        "Tax ID / EIN",
        "Teléfono, correo, dirección",
        "Pie de ticket / reportes",
        "Impuesto por defecto (al crear productos)",
        "Moneda (USD por defecto; MXN opcional)",
    ])

    # 12
    add_heading_styled(doc, "12. Usuarios, roles y auditoría", 1)

    add_heading_styled(doc, "12.1 Usuarios", 2)
    add_steps(doc, [
        "Usuarios → Nuevo → nombre, correo, contraseña, rol y estado activo.",
        "No puedes eliminar tu propia cuenta.",
        "Desactiva usuarios que ya no deban entrar en lugar de compartir contraseñas.",
    ])

    add_heading_styled(doc, "12.2 Roles", 2)
    add_para(doc, "Define conjuntos de permisos (ver productos, crear ventas, cancelar, etc.). El rol «admin» está protegido y no se puede eliminar.")

    add_heading_styled(doc, "12.3 Auditoría", 2)
    add_para(doc, "Registro de cambios importantes (quién hizo qué y cuándo). Útil para control interno y soporte.")

    # 13
    add_heading_styled(doc, "13. Flujos recomendados del día a día", 1)

    add_heading_styled(doc, "13.1 Puesta en marcha (una sola vez)", 2)
    add_steps(doc, [
        "Configura datos de la empresa y moneda USD.",
        "Crea impuestos, unidades, categorías y marcas.",
        "Carga productos con stock inicial o mediante una compra + recepción.",
        "Crea usuarios por rol (cajero, almacén, gerente).",
        "Da de alta proveedores y clientes frecuentes.",
    ])

    add_heading_styled(doc, "13.2 Rutina de ventas (cajero)", 2)
    add_steps(doc, [
        "Abrir caja con el efectivo inicial.",
        "Vender en POS durante el turno.",
        "Consultar ventas si un cliente pide el folio.",
        "Cerrar caja al final del turno con el efectivo contado.",
    ])

    add_heading_styled(doc, "13.3 Rutina de compras / almacén", 2)
    add_steps(doc, [
        "Crear orden de compra al proveedor.",
        "Al llegar la mercancía, recibir (total o parcial).",
        "Verificar kardex y alertas de bajo stock en Dashboard.",
        "Ajustar inventario solo con justificación (merma, conteo físico).",
    ])

    add_heading_styled(doc, "13.4 Cierre de periodo (gerente)", 2)
    add_steps(doc, [
        "Revisar Dashboard del mes.",
        "Exportar reportes de ventas, compras e inventario.",
        "Revisar auditoría si hubo cancelaciones o ajustes atípicos.",
    ])

    # 14
    add_heading_styled(doc, "14. Preguntas frecuentes y problemas comunes", 1)
    add_table(doc, ["Situación", "Qué hacer"], [
        ["No puedo vender / error de caja", "Abre una sesión de caja en Punto de venta."],
        ["No aparece «Cancelar venta»", "Tu rol no tiene permiso o la venta no está completada."],
        ["No puedo cancelar una compra", "Ya tiene recepciones; solo se cancelan órdenes sin recibir."],
        ["Stock no sube al comprar", "Debes «Recibir» la compra; crear la orden no mueve stock."],
        ["No puedo editar el stock en el producto", "Usa Ajustar, Compra/Recepción o una venta; es intencional."],
        ["Lista vacía o error rojo", "Recarga la página; si persiste, avisa a soporte con la pantalla/consola."],
        ["Montos raros / moneda", "Revisa Configuración → Moneda (debe ser USD)."],
        ["Modal no cierra / botones raros", "Recarga el frontend (F5). Usa «Cancelar» o la X del modal."],
    ])

    add_heading_styled(doc, "Glosario rápido", 1)
    add_table(doc, ["Término", "Significado"], [
        ["SKU", "Código interno único del producto"],
        ["Kardex", "Historial de movimientos de inventario"],
        ["POS", "Punto de venta / caja"],
        ["Folio", "Número de documento (C-… compras, V-… ventas)"],
        ["Tax ID / EIN", "Identificación fiscal de la empresa o tercero"],
        ["USD", "Dólares estadounidenses"],
    ])

    add_para(doc, "")
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run("Fin del manual — Inventory System")
    set_run_font(r, size=11, bold=True, color=RGBColor(0x0F, 0x76, 0x6E))

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("URL de acceso local: http://localhost:5173  ·  API: http://localhost:8000")
    set_run_font(r, size=9, color=RGBColor(0x64, 0x74, 0x8B))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
