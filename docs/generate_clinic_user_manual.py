#!/usr/bin/env python3
"""Generate Clinic System user manual (.docx)."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(r"C:\xampp\htdocs\inventory-system\docs\Manual_Usuario_Clinic_System.docx")


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            color=RGBColor(0x0F, 0x76, 0x6E),
        )
    return p


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)
        p.paragraph_format.space_after = Pt(2)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    run = title.add_run("Manual de Usuario — Clinic System")
    set_run_font(run, size=22, bold=True, color=RGBColor(0x0F, 0x76, 0x6E))
    add_para(doc, "Guía breve de módulos clínicos del frontend.")

    add_heading_styled(doc, "1. Acceso", 1)
    add_para(doc, "Inicie sesión con su correo y contraseña. El acceso depende de roles y permisos.")

    add_heading_styled(doc, "2. Dashboard", 1)
    add_bullets(
        doc,
        [
            "Pacientes activos, citas de hoy, consultas del periodo y total de pagos.",
            "Gráfica de citas (últimos 7 días) y listado de próximas citas.",
        ],
    )

    add_heading_styled(doc, "3. Pacientes", 1)
    add_bullets(
        doc,
        [
            "Alta, edición y baja de pacientes.",
            "Ficha con pestañas: Datos, Contactos (CRUD) e Historial (citas, consultas, recetas).",
        ],
    )

    add_heading_styled(doc, "4. Médicos", 1)
    add_bullets(
        doc,
        [
            "Crear médico con nombre, correo, contraseña y cédula.",
            "Asignar especialidades y editar horarios semanales (día 0–6, inicio, fin, minutos por slot).",
        ],
    )

    add_heading_styled(doc, "5. Citas", 1)
    add_bullets(
        doc,
        [
            "Filtrar por fechas y médico.",
            "Crear cita (paciente, médico, especialidad, inicio/fin, motivo).",
            "Cancelar con motivo o reprogramar.",
        ],
    )

    add_heading_styled(doc, "6. Consultas", 1)
    add_bullets(
        doc,
        [
            "Registrar motivo, diagnóstico, tratamiento y observaciones.",
            "Capturar signos vitales y finalizar la consulta.",
            "Desde la ficha: crear receta y registrar pago opcional.",
        ],
    )

    add_heading_styled(doc, "7. Recetas", 1)
    add_bullets(
        doc,
        [
            "Consultar ítems prescritos.",
            "Cancelar receta o abrir PDF.",
        ],
    )

    add_heading_styled(doc, "8. Catálogos y configuración", 1)
    add_bullets(
        doc,
        [
            "Especialidades y medicamentos (CRUD).",
            "Configuración: nombre de clínica, RFC, dirección, teléfono, correo, moneda, duración de cita, pie de ticket y exigir pago para finalizar.",
        ],
    )

    add_heading_styled(doc, "9. Reportes", 1)
    add_para(
        doc,
        "Reportes de pacientes, citas, consultas, médicos y pagos, con export PDF/CSV según permisos.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
